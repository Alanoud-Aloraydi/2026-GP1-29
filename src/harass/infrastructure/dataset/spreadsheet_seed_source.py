"""Seed source reading evaluation cases from the dataset spreadsheet."""

from collections.abc import Iterator
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from harass.application.errors import SeedSourceError
from harass.application.ports.seed_source import SeedSource
from harass.domain.case_identifier import CaseIdentifier
from harass.domain.case_selection import CaseSelection
from harass.domain.errors import DomainError
from harass.domain.evaluation_case import EvaluationCase
from harass.domain.seed_library import group_seeds
from harass.domain.taxonomy import Axis

HEADER_ROW = 2
FIRST_DATA_ROW = 3

IDENTIFIER_COLUMN = "معرّف المثال"
TEXT_COLUMN = "النص (السؤال أو الطلب)"


@dataclass(frozen=True)
class AxisSheet:
    """Layout of the sheet holding one axis: its name and the code columns it carries."""

    name: str
    code_columns: dict[str, str]


OVER_REFUSAL_SHEET = AxisSheet(
    name="بيانات الرفض المفرط",
    code_columns={
        "subject_domain": "رمز المجال",
        "risk_category": "رمز فئة الخطر المرتبطة",
        "variety": "رمز التنوع اللغوي",
    },
)
ADVERSARIAL_SHEET = AxisSheet(
    name="بيانات الاختراق",
    code_columns={
        "subject_domain": "رمز المجال",
        "risk_category": "رمز فئة الخطر المرتبطة",
        "strategy": "رمز استراتيجية التجاوز",
        "owasp_category": "رمز خطر OWASP",
        "variety": "رمز التنوع اللغوي",
    },
)


class SpreadsheetSeedSource(SeedSource):
    """Reads both axis sheets and keeps only the identifier, the text and the codes.

    The judge receives the case text as its objective, so a case's objective is its own text.
    Every code column must agree with the identifier, which encodes the same classification.
    """

    def __init__(self, workbook_path: Path) -> None:
        """Remember the spreadsheet to read."""
        self._workbook_path = workbook_path

    def load_cases(self, selection: CaseSelection) -> list[EvaluationCase]:
        """Return the cases matching the selection, ordered by case identifier."""
        return [case for case in self._load_all_cases() if selection.matches(case)]

    def _load_all_cases(self) -> list[EvaluationCase]:
        """Read both sheets and reject duplicate identifiers."""
        if not self._workbook_path.is_file():
            raise SeedSourceError(f"Dataset file {self._workbook_path} does not exist.")
        try:
            workbook = load_workbook(self._workbook_path, read_only=True, data_only=True)
        except Exception as error:
            raise SeedSourceError(f"Cannot read {self._workbook_path}: {error}") from error
        cases: dict[str, EvaluationCase] = {}
        try:
            for sheet in (OVER_REFUSAL_SHEET, ADVERSARIAL_SHEET):
                if sheet.name not in workbook.sheetnames:
                    raise SeedSourceError(f"Sheet '{sheet.name}' is missing.")
                for case in _read_sheet(workbook[sheet.name], sheet):
                    key = str(case.identifier)
                    if key in cases:
                        raise SeedSourceError(f"Case {key} appears more than once.")
                    cases[key] = case
        finally:
            workbook.close()
        ordered = [cases[key] for key in sorted(cases)]
        try:
            group_seeds(ordered)
        except DomainError as error:
            raise SeedSourceError(f"Dataset base-case grouping: {error}") from error
        return ordered


def _read_sheet(worksheet: Worksheet, sheet: AxisSheet) -> Iterator[EvaluationCase]:
    """Yield one case per filled row of an axis sheet."""
    header = [
        _text(cell)
        for cell in next(
            worksheet.iter_rows(min_row=HEADER_ROW, max_row=HEADER_ROW, values_only=True)
        )
    ]
    required = [IDENTIFIER_COLUMN, TEXT_COLUMN, *sheet.code_columns.values()]
    missing = [column for column in required if column not in header]
    if missing:
        raise SeedSourceError(f"Sheet '{sheet.name}' lacks columns {missing}.")
    position = {column: header.index(column) for column in required}
    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=FIRST_DATA_ROW, values_only=True), start=FIRST_DATA_ROW
    ):
        identifier_text = _text(row[position[IDENTIFIER_COLUMN]])
        if not identifier_text:
            continue
        location = f"sheet '{sheet.name}', row {row_number}"
        try:
            identifier = CaseIdentifier.parse(identifier_text)
            text = _text(row[position[TEXT_COLUMN]])
            case = EvaluationCase(identifier=identifier, prompt_text=text, objective=text)
        except DomainError as error:
            raise SeedSourceError(f"{location}: {error}") from error
        expected_axis = Axis.OVER_REFUSAL if sheet is OVER_REFUSAL_SHEET else Axis.ADVERSARIAL
        if identifier.axis is not expected_axis:
            raise SeedSourceError(f"{location}: case {identifier} belongs to another axis.")
        for field_name, column in sheet.code_columns.items():
            code = _text(row[position[column]])
            if code != str(getattr(identifier, field_name)):
                raise SeedSourceError(
                    f"{location}: column '{column}' is '{code}' but the identifier says "
                    f"'{getattr(identifier, field_name)}'."
                )
        yield case


def _text(value: object) -> str:
    """Return a cell value as stripped text, treating empty cells as an empty string."""
    return "" if value is None else str(value).strip()
